"""Write one airline's scenario workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SCEN = {
    1: ('Fuel uplift - invoice matching', 'YES',
        'Invoices app > open the invoice > Validate > Post to S/4.'),
    2: ('Fuel uplift - invoicing price mismatch', 'YES',
        'Open each P-invoice > Validate. A posts (warning only); B is gated and can be bypassed with a reason; C is gated hard.'),
    3: ('Fuel uplift - invoicing (price components)', 'YES, after Derive Price',
        'Pricing > derivePrice(contractId of the CMP contract, effectiveDate = flight date). Then open each C-invoice > Validate and read INV471/INV472 per line.'),
    4: ('Fuel uplift - quantity mismatch', 'YES',
        'Open each Q-invoice > Validate. Compare line quantity with the ticket and the goods receipt.'),
    5: ('Into-plane - invoice matching', 'PARTIAL - defect',
        'Open the ITP invoice > Validate. INV455 shows fuel + into-plane as two charges (passes); INV452 fails HARD on every line - see Findings F1.'),
    6: ('Un-billed uplift', 'YES (report); INV453 not built',
        'Exception Reports > Un-invoiced tickets (or Unbilled Tickets). Filter by station/date.'),
    7: ('Uplift updated after month-end close', 'YES (B6) / after ledger tool (TP)',
        'FS-to-Finance reconciliation > period of the flight vs the next period > "Posting to subsequent period". Ticket page shows created date and FI posting date.'),
    8: ('Fuelling event does not match flight schedule date', 'YES',
        'Open the ticket (fuelled the evening before) and its order; then Validate its invoice and note the line is left with no flight.'),
    9: ('Mapping of fuelling event to exact flight version', 'YES',
        'Flight page > Dispatch Plans: versions, SUPERSEDED/ACTIVE, version gap. Order page > dispatch plan it names.'),
    10: ('Maintenance uplift (uplift to tail)', 'YES - adds code-list value MAINTENANCE',
         'Fuel Orders filtered by category MAINTENANCE; open ticket (no flight); Validate the MNT invoice.'),
    11: ('Defuelling - invoice matching', 'NO - needs schema (proposed sheets)',
         'Not demonstrable until movement type / DEFUEL_SERVICE charge exist. Data is ready in PROPOSED_DEFUEL_* sheets.'),
    12: ('Defuelling to supplier (credit note)', 'PARTIAL - shows the gap',
         'Load the credit note as an invoice and Validate: INV456 HARD on the negative line (no field marks a defuel). Full demo needs document_type CREDIT_NOTE.'),
    13: ('Defuelling to storage location', 'NO - needs schema (proposed sheets)',
         'Needs storage-location master and movement type. Data ready in PROPOSED_DEFUEL_TICKETS / PROPOSED_STORAGE_LOCATIONS.'),
    14: ('Uplift to tail (preflight uplift)', 'YES',
         'Fuel Tickets filtered UNMATCHED; Validate its invoice (INV463); attach the ticket to the waiting order; Validate again.'),
    15: ('Duplicate fuel ticket - should not be posted', 'PARTIAL - no guard built',
         'Fuel Tickets: search the original ticket number - A and B appear twice. Nothing stops them today (EPD410 DESIGNED). C is the control pair.'),
    16: ('Duplicate invoice - one posted, the other in exception', 'YES',
         'Open the duplicate (A, B) > Validate: GATED by INV473/INV455; the original is already POSTED. C (control) validates CLEAR.'),
    17: ('Accounting audit trail / reconciliation (a-h)', 'PARTIAL - S/4 posts the journals',
         'GL accounts load as RECON_GL_ACCOUNTS; expected journal entries per document are in Expected_Journal_Entries. FuelSphere posts PO/GR/invoice to S/4 and reconciles; it does not book journals itself.'),
}

FINDINGS = [
    ('F1', 'DEFECT', 'INV452 prices an INTO_PLANE line against the FUEL order price.',
     'Every into-plane line on the existing B6 agent invoices fails HARD (e.g. 0.0165 vs 1.2966, -98.7%). Scenario 5 cannot pass until INTO_PLANE lines take their reference from the into-plane contract. No into-plane rate is recorded anywhere in FuelSphere today.'),
    ('F2', 'GAP', 'No defuel model.',
     'INV456 says negative lines are valid for a defuel, but no field on FUEL_TICKETS or INVOICE_ITEMS marks one (invoice-checks.js:702), so every negative line raises HARD. Proposed: FUEL_TICKETS.movement_type, defuel_disposition, storage_location, reference_ticket; INVOICES.document_type + reference_invoice; charge_type DEFUEL_SERVICE.'),
    ('F3', 'GAP', 'No credit-note document type.',
     'INVOICES has no document_type; a credit note can only be a negative invoice. Needed for scenarios 12 and 17e/g.'),
    ('F4', 'GAP', 'No storage-location master.',
     'Scenario 13 needs a station storage location (S/4 plant + LGORT) to move defuelled fuel into.'),
    ('F5', 'GAP', 'No duplicate-ticket guard (EPD410 DESIGNED).',
     'A re-sent or re-keyed ticket saves and posts PO+GR. Criteria proposed in scenario 15. INV455 only catches it if a supplier bills it.'),
    ('F6', 'BEHAVIOUR', 'INV473 and INV455 are symmetric.',
     'Validating the ORIGINAL after the duplicate exists flags the original too. The seed therefore carries the originals as already POSTED. A rule "first captured stays payable" is a decision, not built.'),
    ('F7', 'BEHAVIOUR', 'INV455 matches ticket NUMBER across all suppliers.',
     'Two suppliers using the same ticket number would flag each other. Not seeded as a failure; noted for the duplicate criteria.'),
    ('F8', 'GAP', 'INV453 (unbilled at period close) is registered and not implemented.',
     'Scenario 6 is visible on the report only.'),
    ('F9', 'DATA', 'GL account numbers are placeholders except B6 600780/600781.',
     'Every placeholder row says so in notes. Confirm with the functional consultant before loading RECON_GL_ACCOUNTS.'),
    ('F10', 'DATA', 'TAP company code TP00, prices and into-plane rates are ASSUMED.',
     'No TAP contract data was supplied. Every invoice note says so.'),
    ('F12', 'DEFECT', 'An invoice line fuelled the evening before is attached to the wrong flight.',
     'invoice-service.js (~line 251) takes the ticket\'s own flight, else looks up flight_number on the ticket\'s DELIVERY date. It never uses the ORDER\'s flight. A daily flight fuelled the night before (scenario 8) therefore resolves to the previous day\'s instance - measured on B6 PMA-SJU-755500 and TP ATI-AMS-873239. Using order.flight_ID before the date lookup fixes it.'),
    ('F11', 'BEHAVIOUR', 'S/4 s4_fi_posting_date is never returned today (CR-2).',
     'Scenario 7 sets it on the seeded tickets so the reconciliation can place them; live postings will show "unknown" until CR-2.'),
]

LOAD_ORDER = ['T005_COUNTRY', 'CURRENCY_MASTER', 'MASTER_AIRPORTS', 'AIRCRAFT_MASTER', 'AIRCRAFT_REGISTRATIONS',
              'AIRCRAFT_FUEL_TYPES', 'CARRIER_MASTER', 'ORDER_CATEGORY', 'MASTER_SUPPLIERS', 'MASTER_CONTRACTS',
              'CONTRACT_LOCATIONS', 'CONTRACT_PRODUCTS', 'STATION_SUPPLIER_MAP', 'MARKET_INDICES',
              'MARKET_INDEX_VALUES', 'PRICING_FORMULAS', 'FORMULA_COMPONENTS', 'RECON_GL_ACCOUNTS',
              'FLIGHT_SCHEDULE', 'FLIGHT_DISPATCH', 'FUEL_ORDERS', 'FUEL_DELIVERIES', 'FUEL_TICKETS',
              'INVOICES', 'INVOICE_ITEMS', 'FUEL_BURNS', 'ROB_LEDGER', 'FUEL_STOCK_MOVEMENT_LEDGER',
              'FUEL_STOCK_BALANCE', 'FUEL_BURN_HEADER']

PROPOSED_SCHEMA = [
    ('FUEL_TICKETS', 'movement_type', "String(10) enum UPLIFT | DEFUEL, default UPLIFT", '11, 12, 13'),
    ('FUEL_TICKETS', 'defuel_disposition', 'String(20) enum SUPPLIER_SERVICE | RETURN_TO_SUPPLIER | STORAGE | WRITE_OFF', '11, 12, 13, 17g'),
    ('FUEL_TICKETS', 'storage_location', 'String(10) - S/4 LGORT the fuel moved to', '13'),
    ('FUEL_TICKETS', 'reference_ticket', 'Association to FUEL_TICKETS - the uplift a defuel reverses', '11, 12, 13'),
    ('FUEL_DELIVERIES', 'movement_type / defuel_disposition / storage_location', 'as on the ticket', '11, 12, 13'),
    ('INVOICES', 'document_type', 'String(12) enum INVOICE | CREDIT_NOTE, default INVOICE', '12, 17e'),
    ('INVOICES', 'reference_invoice', 'Association to INVOICES - what a credit note credits', '12'),
    ('INVOICE_ITEMS', 'charge_type', 'add member DEFUEL_SERVICE (enum InvoiceChargeType)', '11, 17f'),
    ('STORAGE_LOCATIONS', '(new entity)', 'carrier, station, plant, storage_location, capacity', '13'),
    ('ORDER_CATEGORY', 'MAINTENANCE (row)', 'Code-list row, not schema. Included in the loadable sheets', '10'),
    ('INVOICE_CHECK_REGISTRY / ticket capture', 'EPD410 duplicate ticket', 'Rule + capture/posting guard per the scenario 15 criteria', '15'),
]

HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E78')
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(wb, name, header, rows, widths=None, wrap=False):
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in rows:
        ws.append([r.get(h, '') if isinstance(r, dict) else r[i] for i, h in enumerate(header)])
    ws.freeze_panes = 'A2'
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h) or min(max(len(str(h)) + 2, 12), 40)
        ws.column_dimensions[get_column_letter(i)].width = w
    if wrap:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
    return ws


def write_workbook(path, code, D, refs, tagger, rows, updates, cols):
    wb = Workbook()
    wb.remove(wb.active)
    name = 'JetBlue (B6) 24-26 Sep 2026' if code == 'B6' else 'TAP Air Portugal (TP) 20-22 Oct 2026'
    readme = [
        f'FuelSphere scenario seed: {name}',
        'Generated by tools/scenario-seed/build.py. NOT LOADED: nothing in db/data was changed.',
        '',
    ]
    if code == 'B6':
        readme += [
            'WHAT THIS IS. The JetBlue seed is already in db/data. This workbook is a DIFF over it:',
            '  ADD_<ENTITY> sheets - new rows, every column named as in the existing CSV (append them).',
            '  UPDATES - one row per changed field: entity, ID, column, current value, new value. Apply after the adds.',
            '  PROPOSED_* sheets - rows the current schema cannot hold (defuel, credit note, storage). Do not load until the schema decision in Proposed_Schema is taken.',
            'Where an existing seed row already demonstrates a scenario it is referenced, not copied (scenarios 5, 6, 9A/9B).',
        ]
    else:
        readme += [
            'WHAT THIS IS. Your uploaded TAP workbook stops at schedule and dispatch. This one carries those rows unchanged',
            '(fuel_order_number / fuel_order_ID filled in) plus the commercial layer: carrier, suppliers, contracts, station',
            'maps, orders, deliveries, tickets, fuel invoices and into-plane invoices, then the scenario rows.',
            'Every sheet is a full set of TP rows to APPEND to db/data/fuelsphere-<ENTITY>.csv in the load order below.',
            'ASSUMES THE JETBLUE SEED IS LOADED (airports, countries, A21N, Jet A product, B6 US suppliers are reused).',
            'ORDER_CATEGORY MAINTENANCE and CURRENCY EUR: skip-if-present (the JetBlue workbook adds MAINTENANCE too).',
        ]
    readme += [
        '',
        'LOAD ORDER: ' + ', '.join(e for e in LOAD_ORDER if e in rows),
        'AFTER LOADING: run tools/seed-burns-and-ledger.js, tools/seed-fifo-valuation.js and tools/seed-idr-rule-status.js'
        + (' (TAP has no ledger rows yet)' if code == 'TP' else ' (moved invoice lines invalidate stored rule status)') + ',',
        'then the harness suite one process per harness (CLAUDE.md section 12) - several harnesses count seeded rows.',
        '',
        'START HERE: Scenario_Index - 17 scenarios x 3 references, what was seeded and what the system should do.',
        'Findings - defects and gaps the scenarios exposed (F1 is a live defect on into-plane invoices).',
        'Expected_Journal_Entries - scenario 17 a-h, the postings each referenced document should produce in S/4.',
        'Scenario_Rows - every row this workbook adds or changes, and the scenario it belongs to.',
        '',
        'SYNTHETIC. Prices, rates, ticket numbers, crew names and GL accounts marked PLACEHOLDER are invented for testing.',
        'Validated before writing: every column is a CDS element, mandatory fields filled, enum values legal, foreign keys',
        'resolve against db/data plus this workbook, invoice headers equal the sum of their lines, no ID or number collides.',
    ]
    ws = wb.create_sheet('README')
    for line in readme:
        ws.append([line])
    ws.column_dimensions['A'].width = 140

    idx = []
    for sid in range(1, 18):
        nm, status, demo = SCEN[sid]
        for r in refs.get(sid, []):
            idx.append({'#': sid, 'Scenario': nm, 'Demonstrable today': status, 'Ref': r['ref'],
                        'Flight (number / local date route)': r['flight'], 'Tail': r['tail'], 'Ticket': r['ticket'],
                        'Invoice': r['invoice'], 'Other reference': r['other'], 'What is seeded': r['seeded'],
                        'Expected system result': r['expected'], 'How to demo': demo})
    hdr = ['#', 'Scenario', 'Demonstrable today', 'Ref', 'Flight (number / local date route)', 'Tail', 'Ticket',
           'Invoice', 'Other reference', 'What is seeded', 'Expected system result', 'How to demo']
    sheet(wb, 'Scenario_Index', hdr, idx, wrap=True,
          widths={'#': 4, 'Scenario': 26, 'Demonstrable today': 18, 'Ref': 5, 'Flight (number / local date route)': 26,
                  'Tail': 10, 'Ticket': 20, 'Invoice': 24, 'Other reference': 36, 'What is seeded': 50,
                  'Expected system result': 60, 'How to demo': 45})
    sheet(wb, 'Findings', ['id', 'kind', 'finding', 'detail'],
          [dict(zip(['id', 'kind', 'finding', 'detail'], f)) for f in FINDINGS], wrap=True,
          widths={'finding': 50, 'detail': 120})

    for ent in LOAD_ORDER:
        rs = rows.get(ent)
        if not rs:
            continue
        header = [c for c in D.t[ent].header if c in cols[ent]]
        sheet(wb, (f'ADD_{ent}' if code == 'B6' else ent), header, rs)
    if updates:
        sheet(wb, 'UPDATES', ['entity', 'key_column', 'key', 'column', 'current_value', 'new_value'], updates,
              widths={'key': 38, 'current_value': 40, 'new_value': 40})
    for name, rs in D.proposed.items():
        if rs:
            header = list(rs[0].keys())
            sheet(wb, f'PROPOSED_{name}', header, rs)
    sheet(wb, 'Proposed_Schema', ['entity', 'element', 'definition', 'scenarios'],
          [dict(zip(['entity', 'element', 'definition', 'scenarios'], p)) for p in PROPOSED_SCHEMA], wrap=True,
          widths={'element': 40, 'definition': 80})
    jh = ['scenario', 'source_scenario', 'reference', 'step', 'gl_account', 'account_role', 'account_name',
          'placeholder_account', 'debit', 'credit', 'currency', 'text', 'basis']
    sheet(wb, 'Expected_Journal_Entries', jh, D.journal, widths={'reference': 34, 'account_name': 44, 'text': 50,
                                                                    'basis': 60, 'placeholder_account': 22})
    sheet(wb, 'Scenario_Rows', ['scenario', 'entity', 'key', 'role'], tagger.rows,
          widths={'key': 38, 'role': 60})
    wb.save(path)
