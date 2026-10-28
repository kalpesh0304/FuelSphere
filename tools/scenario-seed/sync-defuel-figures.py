"""Bring the scenario workbooks' defuel figures into line with the seed.

tools/seed-defuel-scenarios.js posts the defuels through DefuelService and
CUTS a proposal the tail cannot spare (DFL432). The workbooks were written
from the proposals, so their litres, kilograms, charges, credits and
expected journal entries for scenarios 11-13 still showed the proposed
figures. This rewrites them from db/data/fuelsphere-DEFUEL.csv.

What it touches, per seeded defuel (keyed by its ticket number):
  Scenario_Index            the litres quoted in "What is seeded"
  PROPOSED_DEFUEL_DELIVERIES delivered_quantity, fob_after_kg, fob_delta_kg
  PROPOSED_DEFUEL_TICKETS   quantity, quantity_kg
  PROPOSED_DEFUEL_INVOICE_ITEMS  service lines: litres and the seeded charge;
                            credit lines: credited litres, the seeded credit and
                            the rate it implies (the value comes from the
                            tail's FIFO layers, not from one uplift price)
  Expected_Journal_Entries  S17e, S17f, S17g amounts and their basis
  any sheet                 an XXX-FS01 storage tank becomes XXX-TEMP, the
                            tank the seed actually uses

A defuel the seed does not carry (TAP scenario 12: expected credits are
USD-only in the credit memo R1) is left as proposed and listed.

usage: python3 tools/scenario-seed/sync-defuel-figures.py [--write]
"""
import csv, glob, re, sys
import openpyxl

WRITE = '--write' in sys.argv
DISPOSED_L = 300          # scenario 12C: litres disposed of, not credited

seed = {r['ticket_number']: r for r in csv.DictReader(open('db/data/fuelsphere-DEFUEL.csv', encoding='utf-8'), delimiter=';')}


def num(v):
    return float(v) if v not in (None, '') else 0.0


def litres(d):
    return round(num(d['quantity_l']))


def credited_l(ticket, d):
    return litres(d) - (DISPOSED_L if num(d['write_off_usd']) > 0 else 0)


def sheet(wb, name):
    ws = wb[name]
    return ws, {c.value: i for i, c in enumerate(ws[1])}


changes = []


def put(ws, row, col, value, label):
    cell = row[col]
    if str(cell.value) != str(value):
        changes.append(f'  {ws.title} r{cell.row} {label}: {cell.value} -> {value}')
        cell.value = value


for path in sorted(glob.glob('docs/test-data/scenarios/*_Scenario_Seed_*.xlsx')):
    wb = openpyxl.load_workbook(path)
    changes.append(path.split('/')[-1])
    proposed = {}   # ticket -> proposed litres, from the workbook itself

    ws, c = sheet(wb, 'PROPOSED_DEFUEL_TICKETS')
    tick_by_delivery = {}
    for r in ws.iter_rows(min_row=2):
        t = r[c['ticket_number']].value
        tick_by_delivery[r[c['delivery_ID']].value] = t
        proposed[t] = abs(round(num(r[c['quantity']].value)))
        d = seed.get(t)
        if not d:
            changes.append(f'  NOT SEEDED, left as proposed: {t}')
            continue
        put(ws, r, c['quantity'], f'{-litres(d):.2f}', f'{t} quantity')
        put(ws, r, c['quantity_kg'], f'{-num(d["quantity_kg"]):.2f}', f'{t} quantity_kg')

    ws, c = sheet(wb, 'PROPOSED_DEFUEL_DELIVERIES')
    for r in ws.iter_rows(min_row=2):
        t = tick_by_delivery.get(r[c['ID']].value)
        d = seed.get(t)
        if not d:
            continue
        kg = num(d['quantity_kg'])
        put(ws, r, c['delivered_quantity'], f'{-litres(d):.2f}', f'{t} delivered_quantity')
        put(ws, r, c['fob_delta_kg'], f'{-kg:.2f}', f'{t} fob_delta_kg')
        put(ws, r, c['fob_after_kg'], f'{num(r[c["fob_before_kg"]].value) - kg:.2f}', f'{t} fob_after_kg')

    ws, c = sheet(wb, 'PROPOSED_DEFUEL_INVOICE_ITEMS')
    for r in ws.iter_rows(min_row=2):
        t = tick_by_delivery.get(r[c['delivery_ID']].value)
        d = seed.get(t)
        if not d:
            continue
        if num(r[c['quantity']].value) >= 0:            # service charge
            put(ws, r, c['quantity'], f'{litres(d):.3f}', f'{t} service litres')
            put(ws, r, c['net_amount'], f'{num(d["defuel_charges_usd"]):.2f}', f'{t} service charge')
        else:                                            # credit for fuel returned
            cl, credit = credited_l(t, d), num(d['expected_credit_usd'])
            put(ws, r, c['quantity'], f'{-cl:.3f}', f'{t} credited litres')
            put(ws, r, c['unit_price'], f'{credit / cl:.4f}', f'{t} credit rate')
            put(ws, r, c['net_amount'], f'{-credit:.2f}', f'{t} credit')

    ws, c = sheet(wb, 'Scenario_Index')
    col = c['What is seeded']
    for r in ws.iter_rows(min_row=2):
        text = str(r[col].value or '')
        t = next((x for x in seed if x in text), None)
        if not t:
            continue
        d, was = seed[t], proposed.get(t)
        if 'at the uplift price' in text:
            # The credit is valued from the tail's FIFO layers, not one price.
            text = text.replace('at the uplift price', "at the value drawn from the tail's FIFO layers")
            put(ws, r, col, text, f'{t} credit basis')
        if was and was != litres(d):
            new = text.replace(f'(-{was} L)', f'(-{litres(d)} L)').replace(f'{was} L x', f'{litres(d)} L x')
            new = new.replace(f'for -{was} L', f'for -{credited_l(t, d)} L')
            new += f' [Seeded at {litres(d)} L, cut from the proposed {was} L: the tail could not spare more (DFL432).]'
            put(ws, r, col, new, f'{t} text')

    ws, c = sheet(wb, 'Expected_Journal_Entries')
    for r in ws.iter_rows(min_row=2):
        sc = r[c['scenario']].value
        if sc not in ('S17e', 'S17f', 'S17g'):
            continue
        ref = str(r[c['reference']].value).split(' ')[0]
        t = next((x for x in seed if re.sub(r'^DF-(.+)-(\d{6})\d{2}([A-Z])$', r'\1-\2\3', x)
                  == re.sub(r'^(CN|DFS)-', '', ref)), None)
        d = seed.get(t)
        if not d:
            continue
        role, amt = r[c['account_role']].value, ('debit' if r[c['debit']].value not in (None, '') else 'credit')
        L, cl = litres(d), credited_l(t, d)
        value, credit, off = num(d['fuel_value_usd']), num(d['expected_credit_usd']), num(d['write_off_usd'])
        if sc == 'S17f':
            figure, basis = num(d['defuel_charges_usd']), f'{L} L x {num(d["defuel_charges_usd"]) / L:.3f}/L'
        elif sc == 'S17e':
            figure, basis = credit, f'{L} L returned, valued from the tail\'s FIFO layers at {credit:.2f}'
        else:
            figure = {'EXPECTED_CREDIT': credit, 'FUEL_WRITE_OFF': off, 'FUEL_EXP_ADJ': value}[role]
            basis = f'{L} L defuelled worth {value:.2f} (FIFO layers); {cl} L credited, {L - cl} L disposed'
        put(ws, r, c[amt], f'{figure:.2f}', f'{sc} {ref} {role}')
        if not str(r[c['basis']].value).startswith(('Credit note receipt', 'Service invoice receipt')):
            put(ws, r, c['basis'], basis, f'{sc} {ref} basis')

    # The seed stores storage defuels in each station's TEMP tank
    # (tools/seed-storage-locations.py); the proposals named an FS01 tank.
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and re.search(r'\b[A-Z]{3}-FS01\b', cell.value):
                    new = re.sub(r'\b([A-Z]{3})-FS01\b', r'\1-TEMP', cell.value)
                    changes.append(f'  {ws.title} r{cell.row} storage location: ...{cell.value[:0]}FS01 -> TEMP')
                    cell.value = new

    if WRITE:
        wb.save(path)

print('\n'.join(changes))
print('\nwritten' if WRITE else '\nDry run. Re-run with --write to save.')
